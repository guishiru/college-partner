"""Word report generation for completed analysis tasks."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt


class ReportService:
    def build_word_report(
        self,
        output_path: str | Path,
        *,
        filename: str,
        quality_report: dict[str, Any],
        plan: dict[str, Any],
        results: list[dict[str, Any]],
        errors: list[dict[str, Any]],
    ) -> Path:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        document.add_heading("数据分析师分析报告", level=0)
        document.add_paragraph(f"数据文件：{filename}")

        document.add_heading("数据质量摘要", level=1)
        document.add_paragraph(
            f"原始数据：{quality_report.get('original_rows', 0)} 行 × "
            f"{quality_report.get('original_columns', 0)} 列；"
            f"分析视图：{quality_report.get('prepared_rows', 0)} 行 × "
            f"{quality_report.get('prepared_columns', 0)} 列。"
        )
        for change in quality_report.get("changes", []):
            document.add_paragraph(
                f"处理记录：{change.get('message')}（{change.get('count', 0)}处）"
            )
        for issue in quality_report.get("issues", []):
            document.add_paragraph(
                f"数据问题：{issue.get('message')}",
                style="Intense Quote",
            )

        document.add_heading("分析需求", level=1)
        document.add_paragraph(str(plan.get("methods") or "未记录"))

        for result in results:
            document.add_heading(result.get("label", result.get("method", "分析结果")), level=1)
            meta = result.get("meta") or {}
            document.add_paragraph(
                f"方法版本：{meta.get('method_version', '未记录')}；"
                f"有效样本量：{meta.get('analysis_sample_size', meta.get('n', '未记录'))}"
            )
            if result.get("interpretation"):
                document.add_paragraph("文字说明：" + result["interpretation"])
            for table in result.get("tables", []):
                document.add_heading(table.get("title", "结果表"), level=2)
                self._add_dataframe(document, table["data"])

        if errors:
            document.add_heading("未完成项目", level=1)
            for error in errors:
                document.add_paragraph(
                    f"{error.get('label', error.get('method'))}：{error.get('message')}"
                )

        document.save(path)
        return path

    @staticmethod
    def _add_dataframe(document: Document, frame) -> None:
        rows, columns = frame.shape
        table = document.add_table(rows=rows + 1, cols=max(columns, 1))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for index, column in enumerate(frame.columns):
            table.cell(0, index).text = str(column)
        for row_index, (_, row) in enumerate(frame.iterrows(), start=1):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = str(value)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
